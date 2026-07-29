"""Rebuild/extend data/fii_dii_cash.csv from the SEBI FII transaction
disclosure workbooks in the chityy-archive release (owner F6 order
2026-07-27: "when the file lands, commit it and never ask again";
files confirmed FII/DII by owner 2026-07-29).

Source format (one xlsx per month, e.g. Mar_2025.zip -> mar-2025.xlsx):
per-trade rows with TR_DATE, TR_TYPE(*), VALUE (in Rs), RFDE_INSTR_TYPE.
Code meanings verified on Mar-2025 (438k rows): TR_TYPE 01 = buy
(227,564 rows) and 04 = sell (208,997 rows) — the only two codes with
transaction-scale counts; everything else (03/15/07/...) is corporate-
action / off-market noise and is ignored. Equity trades only
(REG_DL_INSTR_EQ). Daily net = (sum buy - sum sell) / 1e7  (Rs -> Rs crore).

Sanity check that validated the mapping: March 2025 output shows heavy
selling 03-17 Mar (-0.3k..-4.7k Cr/day) flipping to strong buying
20-27 Mar (+3.3k..+10.8k Cr/day) — exactly the published FPI story of
that month, at published gross magnitudes (~Rs 3.1L Cr/side).

DII does NOT exist in these files (they are SEBI *FII* disclosures) —
the dii_cash_net column is left EMPTY for derived rows, never invented.
No engine path consumes DII (only fii_flow_score feeds scoring/regime).

Usage:
    python scripts/build_fii_cash_from_disclosures.py <month1.xlsx> [...]
Prints date,fii_cash_net rows (Rs crore) to stdout, sorted by date.
"""

from __future__ import annotations

import sys
from collections import defaultdict

import openpyxl

BUY, SELL = "01", "04"
EQUITY = "REG_DL_INSTR_EQ"


def daily_fii_net(xlsx_path: str) -> dict[str, float]:
    """{iso_date: fii_net_in_crore} for one monthly workbook."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(c).strip() if c else "" for c in next(rows)]
    i_dt = hdr.index("TR_DATE")
    i_ty = hdr.index("TR_TYPE(*)")
    i_val = hdr.index("VALUE (in Rs)")
    i_in = hdr.index("RFDE_INSTR_TYPE")

    net: dict[str, float] = defaultdict(float)
    for r in rows:
        ty = str(r[i_ty] or "").strip()
        if ty not in (BUY, SELL):
            continue
        if str(r[i_in] or "").strip() != EQUITY:
            continue
        try:
            v = float(r[i_val] or 0)
            d = r[i_dt].date().isoformat()
        except (AttributeError, TypeError, ValueError):
            continue
        net[d] += v if ty == BUY else -v
    return {d: round(net[d] / 1e7, 2) for d in sorted(net)}


def main(paths: list[str]) -> None:
    merged: dict[str, float] = {}
    for p in paths:
        month = daily_fii_net(p)
        print(f"# {p}: {len(month)} trade days, "
              f"net {sum(month.values()):+,.0f} Cr", file=sys.stderr)
        merged.update(month)
    print("date,fii_cash_net")
    for d in sorted(merged):
        print(f"{d},{merged[d]}")


if __name__ == "__main__":
    main(sys.argv[1:])
